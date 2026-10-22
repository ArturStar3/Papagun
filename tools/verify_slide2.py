"""End-to-end and failure checks without printing financial or patient data."""
import hashlib
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from io import BytesIO

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import pandas as pd
from openpyxl import load_workbook
from pptx import Presentation
from slide2_export import TEMPLATE, CENTERS, build_slide2, patient_model, _percent, _integer


def digest(path):
    return hashlib.sha256(path.read_bytes()).digest()


def must_fail(exc, action):
    try:
        action()
    except exc:
        return
    raise AssertionError(f'Expected {exc.__name__}')


if __name__ == '__main__':
    sources = [TEMPLATE, CENTERS, *sorted((ROOT / 'Исходники').glob('NL_*.xlsx'))]
    before = {p: digest(p) for p in sources}
    nb = json.loads((ROOT / 'excel_loading.ipynb').read_text(encoding='utf-8'))
    scope = {}
    for cell in nb['cells']:
        if cell['cell_type'] == 'code' and cell.get('metadata', {}).get('papagun_stage') != 'presentation':
            exec(compile(''.join(cell['source']), 'excel_loading.ipynb', 'exec'), scope)
    df, prov = scope['quantity_df'], scope['quantity_provenance']
    original = Presentation(TEMPLATE)
    original_chart = next(s.chart for s in original.slides[1].shapes if s.has_chart)
    labels = [c.label for c in original_chart.plots[0].categories]
    model = patient_model(df, prov, labels, month='сентябрь')
    assert len(model) == 18
    assert len({r['branch'] for r in model}) == 18
    assert all(a['ratio'] >= b['ratio'] for a, b in zip(model, model[1:]))
    # Independent source read: AO/AQ = approved/actual 1C in September.
    wb = load_workbook(ROOT / 'Исходники/NL_Кол. Бюджет.xlsx', read_only=True, data_only=True)
    try:
        for r in model:
            cells = list(wb[r['branch']].iter_rows(min_row=14, max_row=14, max_col=43, values_only=True))[0]
            assert r['plan'] == cells[40] and r['actual'] == cells[42]
    finally:
        wb.close()
    idx = model[0]['source_index']
    for value in (pd.NA, '#DIV/0!', -1, 1.5):
        broken = df.copy()
        broken.at[idx, 'Факт по 1С'] = value
        must_fail(ValueError, lambda: patient_model(broken, prov, labels, month='сентябрь'))
    broken = df.copy()
    broken.at[idx, 'Утверждено'] = 0
    must_fail(ValueError, lambda: patient_model(broken, prov, labels, month='сентябрь'))
    # A second article with the same name must not be aggregated accidentally.
    duplicate_name = df.copy()
    other = prov['Строка Excel'].eq(26) & duplicate_name['Месяц'].eq('сентябрь')
    duplicate_name.loc[other, 'Наименование статей'] = 'Количество пациентов'
    duplicate_name.loc[other, 'Факт по 1С'] = 999999
    assert patient_model(duplicate_name, prov, labels, month='сентябрь') == model
    must_fail(ValueError, lambda: patient_model(df, prov.iloc[::-1], labels, month='сентябрь'))
    must_fail(ValueError, lambda: patient_model(df, prov, labels, month='неизвестно'))
    temporary = TemporaryDirectory(prefix='papagun-slide2-')
    target = Path(temporary.name) / 'slide2.pptx'
    build_slide2(df, prov, output=target)
    must_fail(FileExistsError, lambda: build_slide2(df, prov, output=target))
    must_fail(ValueError, lambda: build_slide2(df, prov, output=TEMPLATE))
    deck = Presentation(target)
    assert len(deck.slides) == 1
    chart = next(s.chart for s in deck.slides[0].shapes if s.has_chart)
    assert len(chart.series) == 2
    assert [c.label for c in chart.plots[0].categories] == [r['label'] for r in model]
    for i, field in enumerate(('plan', 'actual')):
        assert list(chart.series[i].values) == [float(r[field]) for r in model]
    assert sum(s.name.startswith('Исполнение ') for s in deck.slides[0].shapes) == 18
    text = '\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
    plan, actual = sum(r['plan'] for r in model), sum(r['actual'] for r in model)
    assert _integer(plan) + ' пациентов' in text
    assert _integer(actual) + ' пациентов' in text
    assert _percent(actual / plan) in text
    with ZipFile(target) as package:
        assert len([n for n in package.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')]) == 1
        assert len([n for n in package.namelist() if n.startswith('ppt/charts/chart') and n.endswith('.xml')]) == 1
        embeddings = [n for n in package.namelist() if n.endswith('.xlsx')]
        assert len(embeddings) == 1
        embedded = load_workbook(BytesIO(package.read(embeddings[0])), read_only=True, data_only=True)
        try:
            entries = list(embedded.active.values)
            assert len(entries) == 19
            for entry, row in zip(entries[1:], model):
                assert entry[:3] == (row['label'], int(row['plan']), int(row['actual']))
        finally:
            embedded.close()
    assert all(digest(p) == before[p] for p in sources)
    assert all(c.get('outputs', []) == [] and c.get('execution_count') is None
               for c in nb['cells'] if c['cell_type'] == 'code')
    temporary.cleanup()
    print('PASS: 18 centers; independent source, chart, embedded XLSX, totals, failures, notebook hygiene, SHA-256.')
