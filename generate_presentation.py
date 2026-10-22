"""Загрузка книг NL и генерация слайдов 3–34 по шаблону.

Ноутбук excel_loading.ipynb не импортируется и не изменяется.
Месяцы не суммируются. Пропуски и ошибки Excel не заменяются нулями.
Правила серий диаграмм выведены из шаблона и перечислены в SERIES_ROLES;
подробности и проверка — docs/PRESENTATION_GENERATOR.md.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import hashlib
import re
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import _Relationship
from pptx.opc.packuri import PackURI
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent
OUTPUT_DIR = ROOT / 'output'
SOURCES = ROOT / 'Исходники'
TEMPLATE = SOURCES / 'Презентация_NL_9 мес 2026.pptx'
CENTERS = SOURCES / 'Соответствие названий центров в презентации и эксель Nephrolife.xlsx'
EXCEL_FILES = {
    'financial': SOURCES / 'NL_Фин. Бюджет.xlsx',
    'quantity': SOURCES / 'NL_Кол. Бюджет.xlsx',
}

BRANCH_SHEETS = [
    'Актобе', 'Актобе АМЦ', 'Астана Биржан Сал', 'Астана Сарыарка',
    'Атырау', 'Атырау Нурсая', 'Атырау Нурсая 2', 'Караганда ЦАД',
    'Караганда Строителей', 'Кентау', 'Костанай', 'Лисаковск', 'Рудный',
    'Павлодар Дюсенова', 'Павлодар Павлова', 'Экибастуз', 'Петропавловск',
    'Усть-Каменогорск (NL UK)', 'Шымкент', 'Шымкент Абая',
]
# Подписи центров презентации, которых нет в справочнике. Лист книги
# называется «Усть-Каменогорск (NL UK)», поэтому «УК» читается как его сокращение.
CENTER_LABEL_ALIASES = {'ук': 'Усть-Каменогорск (NL UK)'}
ARTICLE_ROWS = {'financial': (9, 193), 'quantity': (9, 141)}
MONTHS = [
    'январь', 'февраль', 'март', 'апрель', 'май', 'июнь',
    'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь',
]
PERIOD_CANONICAL = [
    '1 квартал', '2 квартал', '3 квартал', '4 квартал',
    '1 полугодие', '9 месяцев', '2026',
]
RESULT_COLUMNS = [
    'Наименование филиала', 'Наименование статей', 'Период',
    'Утверждено', 'Факт по норме', 'Факт по 1С', 'Отклонение',
]
METRIC_HEADERS = ['утвержд.', 'факт по норме', 'факт по 1с']
UNFILLED_TEMPLATE_SLIDES = {4, 30, 31, 32, 33}
LAST_TEMPLATE_SLIDE = 34
FIRST_TEMPLATE_SLIDE = 3

# Роли серий по имени ряда в шаблоне. Проверено по значениям шаблона:
# «Столбец1» = Факт / План (слайд 3 и четвёртый ряд квартальных диаграмм),
# «отк:» = План − Факт. Колонка Excel «Отклонение» в диаграммы не пишется:
# её знак в книгах непостоянен.
SERIES_ROLES = {
    'план': 'plan',
    'факт': 'actual',
    'отк:': 'plan_minus_actual',
    'столбец1': 'ratio',
}

# Явные правила для диаграмм шаблона, которые автопривязка не разрешает.
# Слайд 3: подпись «(Процедуры)» неоднозначна, но значения шаблона совпадают
# с первой строкой «Количество процедур» количественной книги (19 из 19 точек).
TEMPLATE_CHART_RULES = {
    (3, 1): {'book': 'quantity', 'article': 'Количество процедур', 'occurrence': 1},
}

TITLE_PREFIXES = (
    'количественные показатели', 'финансовые показатели', 'финансовый показатели',
    'расходы лс', 'расход имн', 'нагрузка на аппараты', 'общие показатели',
    'спасибо', 'nephrolife',
)
PERCENT_RE = re.compile(r'-?\d+(?:[.,]\d+)?\s*%')
YEAR_RE = re.compile(r'(?<!\d)20\d{2}(?!\d)')
PERIOD_PATTERNS = [
    (re.compile(r'(?<!\d)([1-4])\s*-?\s*(?:й\s*)?квартал'), lambda m: f'{m.group(1)} квартал'),
    (re.compile(r'(?<!\d)1\s*полугод(?:ие)?\.?'), lambda m: '1 полугодие'),
    (re.compile(r'(?<!\d)9\s*мес(?:яцев|яца|\.)?'), lambda m: '9 месяцев'),
]
FINANCIAL_BOOK_RE = re.compile(r'тыс\.?\s*т(?:г|енге)|тенге|финанс')


def require_project_root():
    if not (ROOT / 'excel_loading.ipynb').is_file():
        raise FileNotFoundError('Модуль должен лежать в корне проекта Papagun.')


def header_text(value):
    return ' '.join(str(value).split()).casefold()


def compact_key(text):
    key = re.sub(r'[^а-яёa-z0-9]', '', str(text).casefold().replace('\\n', ''))
    return key.replace('12мкрн', '12мкр')


def text_tokens(text):
    return [token for token in re.split(r'[^а-яёa-z0-9]+', str(text).casefold()) if token]


def canonical_period(value):
    """Период из заголовка Excel или подписи категории; None, если это не период."""
    if isinstance(value, bool):
        return None
    year = getattr(value, 'year', None)
    if isinstance(year, int) and 2000 <= year <= 2100:
        return str(year)
    if isinstance(value, (int, float)) and 2000 <= int(value) <= 2100:
        return str(int(value))
    text = header_text(value)
    if text in MONTHS:
        return text
    if re.fullmatch(r'20\d{2}(?:\.0+)?', text):
        return text.split('.')[0]
    stripped = re.sub(r'^за\s+', '', text)
    stripped = re.sub(r'\s*20\d{2}\s*г?\.?', '', stripped).strip(' .')
    stripped = re.sub(r'\s+', ' ', stripped)
    for pattern, build in PERIOD_PATTERNS:
        match = pattern.fullmatch(stripped)
        if match:
            return build(match)
    return None


def find_period_in_text(text):
    """Период внутри произвольного текста (заголовок слайда). Год сам по себе не период."""
    lowered = header_text(text)
    found = []
    for pattern, build in PERIOD_PATTERNS:
        for match in pattern.finditer(lowered):
            found.append(build(match))
    unique = sorted(set(found))
    if len(unique) == 1:
        return unique[0]
    return None


def available_output(filename):
    path = Path(filename)
    suffix = 2
    candidate = path
    while candidate.exists():
        candidate = path.with_name(f'{path.stem}_{suffix}{path.suffix}')
        suffix += 1
    return candidate


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).digest()


def source_hashes():
    paths = [TEMPLATE, CENTERS, *EXCEL_FILES.values()]
    return {path: digest(path) for path in paths}


def assert_sources_unchanged(before):
    after = source_hashes()
    if after != before:
        raise AssertionError('Исходные Excel, PPTX или справочник изменились.')


def chart_number(value):
    if value is None:
        return None
    number = Decimal(str(value))
    if not number.is_finite():
        raise ValueError('Показатель не является конечным числом.')
    if number == number.to_integral_value():
        return int(number)
    return float(number)


@dataclass
class Diagnostic:
    slide: int | None
    chart: int | None
    reason: str
    article_rows: tuple[int, ...] = ()
    period: str | None = None
    branch: str | None = None
    label: str | None = None


@dataclass
class LoadedBooks:
    frames: dict
    provenance: dict
    diagnostics: dict
    articles: dict
    center_map: dict
    workbooks: dict


@dataclass
class ChartFill:
    labels: list
    plan: list
    actual: list
    skipped: list = field(default_factory=list)

    def series(self, role):
        if role == 'plan':
            return list(self.plan)
        if role == 'actual':
            return list(self.actual)
        if role == 'plan_minus_actual':
            return [
                None if plan is None or actual is None else plan - actual
                for plan, actual in zip(self.plan, self.actual)
            ]
        if role == 'ratio':
            return [
                None if plan is None or actual is None or plan == 0 else actual / plan
                for plan, actual in zip(self.plan, self.actual)
            ]
        raise ValueError(f'Неизвестная роль серии: {role}')

    def totals(self):
        pairs = [(plan, actual) for plan, actual in zip(self.plan, self.actual) if plan is not None and actual is not None]
        plan = sum((item[0] for item in pairs), Decimal(0))
        actual = sum((item[1] for item in pairs), Decimal(0))
        ratio = actual / plan if plan != 0 else None
        return plan, actual, actual - plan, ratio


def load_center_map(path=CENTERS):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        mapping = {}
        for title, branch in list(workbook.active.values)[1:]:
            if not title or not branch:
                continue
            if branch == 'Усть-Каменогорск':
                branch = 'Усть-Каменогорск (NL UK)'
            key = compact_key(title)
            if key in mapping:
                raise ValueError('Повторяющееся название в справочнике центров.')
            mapping[key] = branch
        for branch in BRANCH_SHEETS:
            mapping.setdefault(compact_key(branch), branch)
        for key, branch in CENTER_LABEL_ALIASES.items():
            mapping.setdefault(compact_key(key), branch)
        return mapping
    finally:
        workbook.close()


def map_center(label, center_map):
    branch = center_map.get(compact_key(label))
    if branch is None or branch not in BRANCH_SHEETS:
        return None
    return branch


def collect_period_columns(frame):
    columns = {}
    for column, value in enumerate(frame.iloc[5]):
        period = canonical_period(value)
        if period is None:
            continue
        if period in columns:
            raise ValueError(f'Повторяющийся заголовок периода: {period}')
        actual = [header_text(frame.iat[6, column + offset]) for offset in range(3)]
        if actual != METRIC_HEADERS:
            raise ValueError(f'Неожиданные показатели периода: {period}')
        columns[period] = column
    missing_months = [month for month in MONTHS if month not in columns]
    if missing_months:
        raise ValueError('Ожидаются все 12 месяцев.')
    missing_periods = [period for period in PERIOD_CANONICAL if period not in columns]
    if missing_periods:
        raise ValueError('Не найдены блоки периодов: ' + ', '.join(missing_periods))
    return columns


def empty_or_missing(value):
    return pd.isna(value) or value == '' or (isinstance(value, str) and not value.strip())


def build_book_tables(source_book, sheets, path):
    if source_book not in ARTICLE_ROWS:
        raise ValueError("Источник должен быть 'financial' или 'quantity'.")
    missing = [sheet for sheet in BRANCH_SHEETS if sheet not in sheets]
    if missing:
        raise ValueError('Не найдены листы филиалов: ' + ', '.join(missing))

    first_row, last_row = ARTICLE_ROWS[source_book]
    records, origins, issues = [], [], []
    cached_book = load_workbook(path, read_only=True, data_only=True)
    try:
        for branch in BRANCH_SHEETS:
            frame = sheets[branch]
            period_columns = collect_period_columns(frame)
            needed = max(period_columns.values()) + 4
            if frame.shape[0] < last_row or frame.shape[1] < needed:
                raise ValueError(f'Недостаточный размер основной таблицы: {branch}')
            if header_text(frame.iat[5, 2]) != 'наименование статей':
                raise ValueError(f'Не найден заголовок статей в C6: {branch}')
            cached_rows = cached_book[branch].iter_rows(
                min_row=first_row, max_row=last_row, max_col=frame.shape[1],
            )
            for excel_row, cells in enumerate(cached_rows, start=first_row):
                article = frame.iat[excel_row - 1, 2]
                if cells[2].data_type == 'e':
                    raise ValueError(f'Ошибка Excel вместо названия статьи: {branch}, C{excel_row}')
                if empty_or_missing(article):
                    continue
                if not isinstance(article, str):
                    raise ValueError(f'Название статьи не является текстом: {branch}, C{excel_row}')
                for period, column in period_columns.items():
                    values = []
                    coordinates = []
                    for offset, metric in enumerate(RESULT_COLUMNS[3:]):
                        coordinate = f'{get_column_letter(column + offset + 1)}{excel_row}'
                        coordinates.append(coordinate)
                        cell = cells[column + offset]
                        value = frame.iat[excel_row - 1, column + offset]
                        if cell.data_type == 'e':
                            value = cell.value
                            issues.append({
                                'Строка DataFrame': len(records),
                                'Филиал': branch,
                                'Период': period,
                                'Показатель': metric,
                                'Ячейка': coordinate,
                                'Ошибка Excel': value,
                            })
                        elif empty_or_missing(value):
                            value = pd.NA
                        values.append(value)
                    records.append([branch, article, period, *values])
                    origins.append([source_book, branch, excel_row, *coordinates])
    finally:
        cached_book.close()

    result = pd.DataFrame(records, columns=RESULT_COLUMNS, dtype=object)
    provenance = pd.DataFrame(origins, columns=[
        'Книга', 'Лист', 'Строка Excel',
        'Ячейка Утверждено', 'Ячейка Факт по норме', 'Ячейка Факт по 1С', 'Ячейка Отклонение',
    ])
    diagnostics = pd.DataFrame(issues, columns=[
        'Строка DataFrame', 'Филиал', 'Период', 'Показатель', 'Ячейка', 'Ошибка Excel',
    ])
    return result, provenance, diagnostics


def article_catalog(frame, provenance):
    catalog = {}
    for index in frame.index:
        name = frame.at[index, 'Наименование статей']
        row = int(provenance.at[index, 'Строка Excel'])
        catalog.setdefault(name, set()).add(row)
    return {name: tuple(sorted(rows)) for name, rows in catalog.items()}


def load_books():
    require_project_root()
    missing = [key for key, path in EXCEL_FILES.items() if not path.is_file()]
    if missing:
        raise FileNotFoundError('Не найдены книги: ' + ', '.join(missing))
    if not TEMPLATE.is_file() or not CENTERS.is_file():
        raise FileNotFoundError('Не найдены шаблон презентации или справочник центров.')
    workbooks = {
        key: pd.read_excel(
            path, sheet_name=None, header=None, dtype=object,
            keep_default_na=False, engine='openpyxl',
        )
        for key, path in EXCEL_FILES.items()
    }
    frames, provenance, diagnostics, articles = {}, {}, {}, {}
    for key in ('quantity', 'financial'):
        frames[key], provenance[key], diagnostics[key] = build_book_tables(
            key, workbooks[key], EXCEL_FILES[key],
        )
        articles[key] = article_catalog(frames[key], provenance[key])
    return LoadedBooks(
        frames=frames,
        provenance=provenance,
        diagnostics=diagnostics,
        articles=articles,
        center_map=load_center_map(),
        workbooks=workbooks,
    )


# --- Разбор текста слайда ---------------------------------------------------

def is_percent_text(text):
    return bool(PERCENT_RE.fullmatch(header_text(text)))


def is_title_text(text):
    compact = header_text(text)
    if not compact or is_percent_text(compact):
        return True
    if YEAR_RE.search(compact):
        return True
    if '(количество)' in compact or '(тыс' in compact:
        return True
    return compact.startswith(TITLE_PREFIXES)


def match_article(caption, articles):
    """Единственная строка Excel по подписи. Порядок: точное имя, префикс,
    сокращение по словам («Утилизация м.о.»), затем вхождение не короче 6 знаков."""
    caption_key = compact_key(caption)
    if not caption_key or is_title_text(caption):
        return None, 'подпись не является статьёй'
    caption_tokens = text_tokens(caption)
    exact, prefix, abbreviated, partial = [], [], [], []
    for name, rows in articles.items():
        name_key = compact_key(name)
        if name_key == caption_key:
            exact.append((name, rows))
        elif len(caption_key) >= 3 and name_key.startswith(caption_key):
            prefix.append((name, rows))
        elif tokens_abbreviate(caption_tokens, text_tokens(name)):
            abbreviated.append((name, rows))
        elif min(len(caption_key), len(name_key)) >= 6 and (
            caption_key in name_key or name_key in caption_key
        ):
            partial.append((name, rows))
    chosen = exact or prefix or abbreviated or partial
    if not chosen:
        return None, 'статья не найдена'
    unique_rows = tuple(sorted({row for _name, rows in chosen for row in rows}))
    if len(unique_rows) != 1:
        return unique_rows, 'статья неоднозначна'
    return unique_rows, None


def tokens_abbreviate(short, full):
    if len(short) < 2 or len(short) != len(full) or len(short[0]) < 4:
        return False
    return all(word.startswith(token) for token, word in zip(short, full))


def article_rows_by_rule(rule, articles):
    rows = articles.get(rule['article'])
    if not rows:
        return None, 'статья правила не найдена в книге'
    occurrence = rule.get('occurrence', 1)
    if occurrence > len(rows):
        return None, 'в книге меньше повторов статьи, чем требует правило'
    return (rows[occurrence - 1],), None


def book_for_slide_text(texts):
    blob = ' '.join(header_text(text) for text in texts)
    return 'financial' if FINANCIAL_BOOK_RE.search(blob) else 'quantity'


def classify_categories(categories):
    keys = [header_text(label) for label in categories]
    if len(keys) >= 2 and keys[0].startswith('план') and keys[1].startswith('факт'):
        return 'unsupported'
    periods = [canonical_period(label) for label in categories]
    if periods and all(item is not None for item in periods):
        return 'period'
    return 'branch'


def shape_text(shape):
    return ' '.join(part.text.strip() for part in shape.text_frame.paragraphs if part.text.strip())


def slide_texts(slide):
    return [shape_text(shape) for shape in slide.shapes if shape.has_text_frame and shape_text(shape)]


def slide_article_labels(slide):
    labels = []
    for shape in slide.shapes:
        if not shape.has_text_frame or shape.has_chart:
            continue
        text = shape_text(shape)
        if text and not is_title_text(text):
            labels.append((shape, text))
    return labels


def nearest_article_text(shape, labels, limit=Inches(1.4)):
    best, distance = None, None
    for other, text in labels:
        gap = abs(other.top - shape.top)
        if distance is None or gap < distance:
            best, distance = text, gap
    if distance is None or distance > limit:
        return None
    return best


def period_for_slide(texts, charts):
    """Период филиальных диаграмм: заголовок слайда, иначе единственный период
    одноточечных диаграмм того же слайда."""
    for text in texts:
        period = find_period_in_text(text)
        if period is not None:
            return period, None
    siblings = set()
    for shape in charts:
        categories = [category.label for category in shape.chart.plots[0].categories]
        if len(categories) == 1:
            period = canonical_period(categories[0])
            if period is not None:
                siblings.add(period)
    if len(siblings) == 1:
        return siblings.pop(), 'период взят из соседней одноточечной диаграммы'
    return None, None


# --- Значения ---------------------------------------------------------------

def lookup_metric(loaded, book, branch, article_row, period):
    frame = loaded.frames[book]
    provenance = loaded.provenance[book]
    selected = frame.loc[
        frame['Наименование филиала'].eq(branch)
        & frame['Период'].eq(period)
        & provenance['Строка Excel'].eq(article_row)
    ]
    if len(selected) != 1:
        return None, 'нет единственной строки показателя'
    row = selected.iloc[0]
    origin = provenance.loc[selected.index[0]]
    if origin['Лист'] != branch:
        return None, 'лист происхождения не совпадает с филиалом'
    values = {}
    for column in ('Утверждено', 'Факт по 1С'):
        value = row[column]
        if isinstance(value, str) and value.startswith('#'):
            return None, f'ошибка Excel в показателе {column}'
        if empty_or_missing(value):
            return None, f'пропуск в показателе {column}'
        try:
            number = Decimal(str(value))
        except Exception:
            return None, f'показатель {column} не является числом'
        if not number.is_finite():
            return None, f'показатель {column} не является числом'
        values[column] = number
    return values, None


def collect_points(loaded, book, article_row, pairs, diagnostics, slide, chart_index):
    """Сохраняет исходные подписи категорий. Нет значения — None в ряде, не удаление точки."""
    points, skipped, usable = [], [], 0
    for label, branch, period in pairs:
        if branch is None:
            skipped.append(label)
            points.append((label, None))
            continue
        values, reason = lookup_metric(loaded, book, branch, article_row, period)
        if reason:
            skipped.append(label)
            diagnostics.append(Diagnostic(slide, chart_index, reason, (article_row,), period, branch, label))
            points.append((label, None))
            continue
        if values['Утверждено'] == 0:
            diagnostics.append(Diagnostic(
                slide, chart_index, 'нулевой план: отношение не вычисляется',
                (article_row,), period, branch, label,
            ))
        points.append((label, values))
        usable += 1
    if usable == 0:
        return None
    return ChartFill(
        labels=[item[0] for item in points],
        plan=[item[1]['Утверждено'] if item[1] else None for item in points],
        actual=[item[1]['Факт по 1С'] if item[1] else None for item in points],
        skipped=skipped,
    )


def fill_from_branches(loaded, book, article_row, period, labels, diagnostics, slide, chart_index):
    pairs = []
    for label in labels:
        branch = map_center(label, loaded.center_map)
        if branch is None:
            diagnostics.append(Diagnostic(slide, chart_index, 'центр не сопоставлен', (article_row,), period, None, label))
            pairs.append((label, None, period))
            continue
        pairs.append((label, branch, period))
    return collect_points(loaded, book, article_row, pairs, diagnostics, slide, chart_index)


def fill_from_periods(loaded, book, article_row, periods, branch, diagnostics, slide, chart_index, labels=None):
    labels = list(labels) if labels is not None else list(periods)
    pairs = [(label, branch, period) for label, period in zip(labels, periods)]
    return collect_points(loaded, book, article_row, pairs, diagnostics, slide, chart_index)


# --- Диаграммы и слайды ------------------------------------------------------

def series_roles(chart):
    roles = []
    for series in chart.series:
        role = SERIES_ROLES.get(header_text(series.name))
        if role is None:
            return None
        roles.append(role)
    return roles


def replace_chart(chart, fill, original_count=None):
    roles = series_roles(chart)
    if roles is None:
        raise ValueError('Серия с неизвестной ролью.')
    data = CategoryChartData()
    data.categories = fill.labels
    for series, role in zip(chart.series, roles):
        data.add_series(series.name, [chart_number(value) for value in fill.series(role)])
    chart.replace_data(data)
    if original_count:
        scale_plot_area(chart, len(fill.labels) / original_count)


def scale_plot_area(chart, factor):
    layouts = chart._chartSpace.xpath('.//c:plotArea/c:layout/c:manualLayout')
    if not layouts or factor >= 1:
        return
    width_node = layouts[0].xpath('./c:w')
    x_node = layouts[0].xpath('./c:x')
    if not width_node or not x_node:
        return
    width = float(width_node[0].get('val'))
    left = float(x_node[0].get('val'))
    new_width = width * factor
    width_node[0].set('val', repr(new_width))
    x_node[0].set('val', repr(left + (width - new_width) / 2))


def blank_chart(chart):
    """Снять значения образца, сохранив категории и имена рядов."""
    data = CategoryChartData()
    data.categories = [category.label for category in chart.plots[0].categories]
    count = len(data.categories)
    for series in chart.series:
        data.add_series(series.name, [None] * count)
    chart.replace_data(data)


def drop_slides(deck, keep):
    keep_ids = {slide.slide_id for slide in keep}
    for identifier in list(deck.slides._sldIdLst):
        if int(identifier.id) in keep_ids:
            continue
        deck.part.drop_rel(identifier.rId)
        deck.slides._sldIdLst.remove(identifier)
    for slide in deck.slides:
        for relation in list(slide.part.rels.values()):
            if relation.reltype.endswith('/notesSlide'):
                slide.part.drop_rel(relation.rId)
    for node in deck.part._element.xpath('./p:custShowLst | ./p:extLst'):
        node.getparent().remove(node)


def save_presentation(deck, output):
    output = Path(output)
    if output.suffix.lower() != '.pptx':
        raise ValueError('Выходной файл должен иметь расширение .pptx.')
    if output.resolve().is_relative_to(SOURCES.resolve()):
        raise ValueError('Сохранение в каталог исходников запрещено.')
    if output.exists():
        raise FileExistsError('Результат уже существует; укажите новое имя файла.')
    output.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    deck.save(buffer)
    with output.open('xb') as target:
        target.write(buffer.getvalue())
    return output


def join_slides(paths):
    result = Presentation(paths[0])
    package = result.part.package
    layout = result.slides[0].slide_layout.part
    reserved = {str(part.partname) for part in package.iter_parts()}
    for path in paths[1:]:
        source = Presentation(path)
        copied = {}

        def copy_part(part):
            if part in copied:
                return copied[part]
            pattern = re.sub(r'\d+(?=\.[^.]+$)', '%d', str(part.partname))
            if '%d' not in pattern:
                stem, suffix = str(part.partname).rsplit('.', 1)
                pattern = stem + '%d.' + suffix
            index = 1
            while pattern % index in reserved:
                index += 1
            name = PackURI(pattern % index)
            reserved.add(str(name))
            clone = type(part).load(name, part.content_type, package, part.blob)
            copied[part] = clone
            for relation in part.rels.values():
                if relation.is_external:
                    target, mode = relation.target_ref, 'External'
                else:
                    target = layout if relation.reltype == RT.SLIDE_LAYOUT else copy_part(relation.target_part)
                    mode = 'Internal'
                clone.rels._rels[relation.rId] = _Relationship(
                    name.baseURI, relation.rId, relation.reltype, mode, target,
                )
            return clone

        slide_part = copy_part(source.slides[0].part)
        rel_id = result.part.relate_to(slide_part, RT.SLIDE)
        result.slides._sldIdLst.add_sldId(rel_id)
    return result


SUMMARY_DECIMALS = Decimal('0.01')


def format_amount(value):
    """Число для сводного текста: разряды через пробел, не более двух знаков после
    запятой (только отображение; ряды диаграмм не округляются)."""
    number = Decimal(value)
    if number == number.to_integral_value():
        text = f'{int(number):,}'
    else:
        text = f'{number.quantize(SUMMARY_DECIMALS).normalize():,f}'
    return text.replace(',', ' ').replace('.', ',')


def format_percent(ratio):
    return f'{(ratio * 100).quantize(Decimal("0.1"))}%'.replace('.', ',')


def update_summary(slide, fill, period, unit_word=''):
    """Сводный текст слайда 3: суммы показанных точек. Структура абзацев проверена
    по шаблону: «План:», «Факт:», «Отклонение:», «Исполнение:». Возвращает причину,
    если структура не распознана. fill=None снимает числа образца без расчёта."""
    if fill is None:
        values = dict.fromkeys(('план:', 'факт:', 'отклонение:', 'исполнение:'), '—')
    else:
        plan, actual, deviation, ratio = fill.totals()
        suffix = f' {unit_word}'.rstrip() if unit_word else ''
        values = {
            'план:': format_amount(plan) + suffix,
            'факт:': format_amount(actual) + suffix,
            'отклонение:': format_amount(deviation) + suffix,
            'исполнение:': format_percent(ratio) if ratio is not None else '—',
        }
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        paragraphs = shape.text_frame.paragraphs
        starts = [header_text(p.runs[0].text) if p.runs else '' for p in paragraphs]
        if not all(key in starts for key in values):
            continue
        for paragraph, start in zip(paragraphs, starts):
            if start in values:
                for run in paragraph.runs[2:]:
                    run._r.getparent().remove(run._r)
                if len(paragraph.runs) < 2:
                    paragraph.add_run()
                paragraph.runs[1].text = ' ' + values[start] if not paragraph.runs[0].text.endswith(' ') else values[start]
            elif start.startswith('общие показатели') and find_period_in_text(paragraph.text) != period:
                paragraph.runs[0].text = f'Общие показатели ({period})'
                for run in paragraph.runs[1:]:
                    run._r.getparent().remove(run._r)
        return None
    return 'сводный текст не распознан и не пересчитан'


def set_title(slide, title):
    """Заголовок слайда — самая широкая текстовая фигура, похожая на заголовок."""
    shapes = [
        shape for shape in slide.shapes
        if shape.has_text_frame and shape_text(shape)
        and is_title_text(shape_text(shape)) and not is_percent_text(shape_text(shape))
        and not header_text(shape_text(shape)).startswith('общие показатели')
    ]
    if not shapes:
        return
    frame = max(shapes, key=lambda shape: shape.width).text_frame
    first = frame.paragraphs[0]
    for paragraph in list(frame.paragraphs[1:]):
        paragraph._p.getparent().remove(paragraph._p)
    for run in first.runs[1:]:
        run._r.getparent().remove(run._r)
    if not first.runs:
        first.add_run()
    first.runs[0].text = title


def clear_percent_texts(slide, diagnostics, number):
    for shape in slide.shapes:
        if shape.has_text_frame and is_percent_text(shape_text(shape)):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = ''
            diagnostics.append(Diagnostic(number, None, 'процентная надпись образца очищена: нет правила пересчёта'))


def generate_from_template(loaded, output=None, *, period_branch=None, stale_policy='clear'):
    """Копия слайдов 3–34. Заполняются только однозначные диаграммы.
    stale_policy='clear' снимает значения образца с незаполненных диаграмм
    и очищает процентные надписи; 'keep' оставляет их как в шаблоне."""
    if stale_policy not in ('clear', 'keep'):
        raise ValueError("stale_policy должен быть 'clear' или 'keep'.")
    diagnostics = []
    deck = Presentation(TEMPLATE)
    keep = []
    filled = {}
    for number, slide in enumerate(deck.slides, 1):
        if number < FIRST_TEMPLATE_SLIDE:
            continue
        keep.append(slide)
        charts = [shape for shape in slide.shapes if shape.has_chart]
        filled[number] = set()
        if number == LAST_TEMPLATE_SLIDE or not charts:
            diagnostics.append(Diagnostic(number, None, 'слайд без расчётных диаграмм скопирован'))
            continue
        if number in UNFILLED_TEMPLATE_SLIDES:
            diagnostics.append(Diagnostic(number, None, 'макет не заполняется: нет подтверждённого правила серий'))
        else:
            fill_template_slide(loaded, slide, number, charts, diagnostics, filled[number], period_branch)
        if stale_policy == 'clear':
            for chart_index, shape in enumerate(charts, 1):
                if chart_index not in filled[number]:
                    blank_chart(shape.chart)
                    diagnostics.append(Diagnostic(number, chart_index, 'значения образца сняты с незаполненной диаграммы'))
            clear_percent_texts(slide, diagnostics, number)
        else:
            if any(is_percent_text(text) for text in slide_texts(slide)):
                diagnostics.append(Diagnostic(number, None, 'процентная надпись образца не пересчитана'))
    drop_slides(deck, keep)
    if output is None:
        output = available_output(OUTPUT_DIR / 'NL_slides_3_34.pptx')
    path = save_presentation(deck, output)
    return path, diagnostics


def fill_template_slide(loaded, slide, number, charts, diagnostics, filled, period_branch):
    texts = slide_texts(slide)
    book = book_for_slide_text(texts)
    labels = slide_article_labels(slide)
    slide_period, period_note = period_for_slide(texts, charts)
    for chart_index, shape in enumerate(charts, 1):
        chart = shape.chart
        categories = [category.label for category in chart.plots[0].categories]
        kind = classify_categories(categories)
        if kind == 'unsupported':
            diagnostics.append(Diagnostic(number, chart_index, 'неподдерживаемые категории План/Факт ср.'))
            continue
        if series_roles(chart) is None:
            diagnostics.append(Diagnostic(number, chart_index, 'ряд с неизвестным именем: правило серии отсутствует'))
            continue
        rule = TEMPLATE_CHART_RULES.get((number, chart_index))
        caption = nearest_article_text(shape, labels)
        if rule:
            book_for_chart = rule.get('book', book)
            rows, reason = article_rows_by_rule(rule, loaded.articles[book_for_chart])
            caption = rule['article']
        else:
            book_for_chart = book
            rows, reason = match_article(caption or '', loaded.articles[book])
        if reason:
            diagnostics.append(Diagnostic(number, chart_index, reason, rows or (), None, None, caption))
            continue
        article_row = rows[0]
        if kind == 'period':
            if period_branch is None:
                diagnostics.append(Diagnostic(
                    number, chart_index, 'диаграмма периодов не заполнена: нет правила итога по сети',
                    (article_row,), None, None, caption,
                ))
                continue
            periods = [canonical_period(label) for label in categories]
            diagnostics.append(Diagnostic(
                number, chart_index, 'диаграмма периодов построена по одному филиалу',
                (article_row,), None, period_branch, caption,
            ))
            fill = fill_from_periods(
                loaded, book_for_chart, article_row, periods, period_branch,
                diagnostics, number, chart_index, labels=categories,
            )
        else:
            if slide_period is None:
                diagnostics.append(Diagnostic(number, chart_index, 'период не определён', (article_row,), None, None, caption))
                continue
            if period_note:
                diagnostics.append(Diagnostic(number, chart_index, period_note, (article_row,), slide_period, None, caption))
            fill = fill_from_branches(
                loaded, book_for_chart, article_row, slide_period, categories, diagnostics, number, chart_index,
            )
        if fill is None:
            diagnostics.append(Diagnostic(number, chart_index, 'нет пригодных точек для замены', (article_row,), None, None, caption))
            continue
        scale_from = len(categories) if kind == 'branch' and len(charts) == 1 else None
        replace_chart(chart, fill, scale_from)
        filled.add(chart_index)
        if number == FIRST_TEMPLATE_SLIDE and slide_period:
            unit = 'процедуры' if rule and rule.get('article') == 'Количество процедур' else ''
            summary_reason = update_summary(slide, fill, slide_period, unit)
            if summary_reason:
                diagnostics.append(Diagnostic(number, chart_index, summary_reason))


def article_name(loaded, book, article_row):
    names = loaded.frames[book].loc[
        loaded.provenance[book]['Строка Excel'].eq(article_row),
        'Наименование статей',
    ].unique()
    if len(names) != 1:
        raise ValueError('Не удалось определить название статьи по номеру строки.')
    return names[0]


def one_chart_slide(fill, original_categories, title, period, output, diagnostics, slide_tag, *, summarize=True):
    """Слайд по макету слайда 3. summarize=False — для наборов периодов:
    их суммы не имеют смысла (9 месяцев включает кварталы), числа образца снимаются."""
    deck = Presentation(TEMPLATE)
    slide = deck.slides[FIRST_TEMPLATE_SLIDE - 1]
    drop_slides(deck, [slide])
    chart_shape = next(shape for shape in slide.shapes if shape.has_chart)
    replace_chart(chart_shape.chart, fill, original_categories)
    set_title(slide, title)
    reason = update_summary(slide, fill if summarize else None, period)
    if reason:
        diagnostics.append(Diagnostic(slide_tag, 1, reason))
    elif not summarize:
        diagnostics.append(Diagnostic(slide_tag, 1, 'сводный текст не вычисляется для набора периодов'))
    return save_presentation(deck, output)


def relabel_captions(slide, assignments):
    """Подписи статей: все нешаблонные подписи очищаются, ближайшая к каждой
    диаграмме получает новое имя."""
    labels = slide_article_labels(slide)
    for shape, _text in labels:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.text = ''
    for chart_shape, text in assignments:
        best, distance = None, None
        for shape, _text in labels:
            gap = abs(shape.top - chart_shape.top)
            if distance is None or gap < distance:
                best, distance = shape, gap
        if best is None:
            continue
        paragraph = best.text_frame.paragraphs[0]
        if not paragraph.runs:
            paragraph.add_run()
        paragraph.runs[0].text = text


def generate_variant(loaded, articles, periods, branches, output, *, book='quantity'):
    """Отдельные слайды по явным номерам строк Excel, периодам и филиалам.
    Диаграммы периодов строятся по первому филиалу списка — без итога по сети."""
    diagnostics = []
    if not articles or not periods or not branches:
        raise ValueError('Нужны непустые списки статей, периодов и филиалов.')
    unknown = [branch for branch in branches if branch not in BRANCH_SHEETS]
    if unknown:
        diagnostics.append(Diagnostic(None, None, 'филиал отсутствует в схеме загрузки', (), None, unknown[0]))
        return None, diagnostics
    unknown_periods = [period for period in periods if period not in PERIOD_CANONICAL and period not in MONTHS]
    if unknown_periods:
        diagnostics.append(Diagnostic(None, None, 'период отсутствует в схеме загрузки', (), unknown_periods[0]))
        return None, diagnostics
    paths = []
    with TemporaryDirectory(prefix='papagun-variant-') as temporary:
        folder = Path(temporary)
        template = Presentation(TEMPLATE)
        slide3 = template.slides[FIRST_TEMPLATE_SLIDE - 1]
        sample_categories = len(next(shape.chart.plots[0].categories for shape in slide3.shapes if shape.has_chart))
        for offset, article_row in enumerate(articles, 1):
            name = article_name(loaded, book, article_row)
            fill = fill_from_branches(
                loaded, book, article_row, periods[0], branches, diagnostics, offset, 1,
            )
            if fill is None:
                diagnostics.append(Diagnostic(offset, 1, 'вариант по филиалам не собран', (article_row,), periods[0]))
                continue
            paths.append(one_chart_slide(
                fill, sample_categories, f'{name.strip()} ({periods[0]})', periods[0],
                folder / f'centers_{article_row}.pptx', diagnostics, offset,
            ))
        if len(articles) >= 2:
            deck = Presentation(TEMPLATE)
            target = deck.slides[6]
            drop_slides(deck, [target])
            period_charts = [shape for shape in target.shapes if shape.has_chart]
            assignments = []
            for chart_index, (shape, article_row) in enumerate(zip(period_charts, articles[:2]), 1):
                fill = fill_from_periods(
                    loaded, book, article_row, periods, branches[0], diagnostics, 100, chart_index,
                )
                if fill is None:
                    blank_chart(shape.chart)
                    continue
                replace_chart(shape.chart, fill)
                assignments.append((shape, article_name(loaded, book, article_row).strip()))
            for shape in period_charts[len(articles[:2]):]:
                blank_chart(shape.chart)
            if assignments:
                clear_percent_texts(target, diagnostics, 100)
                relabel_captions(target, assignments)
                set_title(target, f'{branches[0]}: периоды')
                diagnostics.append(Diagnostic(
                    100, None, 'диаграмма периодов построена по одному филиалу',
                    tuple(articles[:2]), None, branches[0],
                ))
                paths.append(save_presentation(deck, folder / 'periods_slide7.pptx'))
        for offset, article_row in enumerate(articles, 1):
            fill = fill_from_periods(
                loaded, book, article_row, periods, branches[0], diagnostics, 200 + offset, 1,
            )
            if fill is None:
                diagnostics.append(Diagnostic(200 + offset, 1, 'вариант по периодам не собран', (article_row,)))
                continue
            diagnostics.append(Diagnostic(
                200 + offset, 1, 'диаграмма периодов построена по одному филиалу',
                (article_row,), None, branches[0],
            ))
            name = article_name(loaded, book, article_row)
            paths.append(one_chart_slide(
                fill, len(periods), f'{name.strip()} ({branches[0]})', ', '.join(periods),
                folder / f'periods_{article_row}.pptx', diagnostics, 200 + offset, summarize=False,
            ))
        if not paths:
            return None, diagnostics
        combined = join_slides(paths)
        buffer = BytesIO()
        combined.save(buffer)
        result = Presentation(BytesIO(buffer.getvalue()))
        return save_presentation(result, Path(output)), diagnostics


# --- Проверки ---------------------------------------------------------------

def independent_cell(path, sheet, coordinate):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return workbook[sheet][coordinate].value
    finally:
        workbook.close()


def verify_fill_against_excel(loaded, book, article_row, period, branch):
    values, reason = lookup_metric(loaded, book, branch, article_row, period)
    if reason:
        return False, reason
    frame = loaded.frames[book]
    provenance = loaded.provenance[book]
    selected = frame.loc[
        frame['Наименование филиала'].eq(branch)
        & frame['Период'].eq(period)
        & provenance['Строка Excel'].eq(article_row)
    ]
    origin = provenance.loc[selected.index[0]]
    mapping = {
        'Утверждено': origin['Ячейка Утверждено'],
        'Факт по 1С': origin['Ячейка Факт по 1С'],
    }
    for column, coordinate in mapping.items():
        cell_value = independent_cell(EXCEL_FILES[book], branch, coordinate)
        if Decimal(str(cell_value)) != values[column]:
            return False, 'расхождение с независимым чтением openpyxl'
    return True, None


def chart_payloads(path):
    deck = Presentation(path)
    payloads = []
    for slide in deck.slides:
        for shape in slide.shapes:
            if not shape.has_chart:
                continue
            chart = shape.chart
            categories = [category.label for category in chart.plots[0].categories]
            series = [(item.name, list(item.values)) for item in chart.series]
            payloads.append((categories, series))
    return payloads


def verify_series_rules(path):
    """Все заполненные ряды подчиняются SERIES_ROLES: отк: = План − Факт, Столбец1 = Факт / План."""
    checked = 0
    for categories, series in chart_payloads(path):
        roles = [SERIES_ROLES.get(header_text(name)) for name, _values in series]
        if None in roles:
            continue
        by_role = {role: values for role, (_name, values) in zip(roles, series)}
        plan, actual = by_role.get('plan'), by_role.get('actual')
        if plan is None or actual is None or all(value is None for value in plan):
            continue
        for index in range(len(categories)):
            if plan[index] is None or actual[index] is None:
                continue
            if 'plan_minus_actual' in by_role:
                expected = plan[index] - actual[index]
                if abs(by_role['plan_minus_actual'][index] - expected) > 1e-6:
                    raise AssertionError('Ряд «отк:» не равен План − Факт.')
            if 'ratio' in by_role:
                got = by_role['ratio'][index]
                if plan[index] == 0:
                    if got is not None:
                        raise AssertionError('Отношение при нулевом плане должно быть пустым.')
                elif abs(got - actual[index] / plan[index]) > 1e-9:
                    raise AssertionError('Ряд «Столбец1» не равен Факт / План.')
            checked += 1
    return checked


def assert_no_values_in_text(text):
    forbidden = ('Утверждено', 'Факт по 1С', 'Отклонение')
    if any(name in text and any(char.isdigit() for char in text) for name in forbidden):
        raise AssertionError('В отчёт попали значения показателей.')


def summarize_diagnostics(items):
    counts = {}
    for item in items:
        counts[item.reason] = counts.get(item.reason, 0) + 1
    return counts


def compare_with_reference(generated, template=TEMPLATE):
    """Структурное сравнение результата со слайдами 3–34 шаблона.
    Возвращает счётчики по слайдам без значений показателей."""
    source = Presentation(template)
    result = Presentation(generated)
    reference_slides = list(source.slides)[FIRST_TEMPLATE_SLIDE - 1:]
    result_slides = list(result.slides)
    report = []
    if len(reference_slides) != len(result_slides):
        raise AssertionError('Число слайдов результата не равно числу слайдов 3–34 шаблона.')
    for offset, (ref, gen) in enumerate(zip(reference_slides, result_slides)):
        number = FIRST_TEMPLATE_SLIDE + offset
        ref_charts = [shape.chart for shape in ref.shapes if shape.has_chart]
        gen_charts = [shape.chart for shape in gen.shapes if shape.has_chart]
        row = {
            'slide': number,
            'shapes_equal': len(list(ref.shapes)) == len(list(gen.shapes)),
            'layout_equal': ref.slide_layout.name == gen.slide_layout.name,
            'charts': len(ref_charts),
            'charts_equal_count': len(ref_charts) == len(gen_charts),
            'structure_equal': 0,
            'values_same_as_reference': 0,
            'values_replaced': 0,
            'values_blank': 0,
            'texts_equal': slide_texts(ref) == slide_texts(gen),
        }
        for ref_chart, gen_chart in zip(ref_charts, gen_charts):
            ref_categories = [category.label for category in ref_chart.plots[0].categories]
            gen_categories = [category.label for category in gen_chart.plots[0].categories]
            ref_names = [series.name for series in ref_chart.series]
            gen_names = [series.name for series in gen_chart.series]
            if ref_categories == gen_categories and ref_names == gen_names:
                row['structure_equal'] += 1
            gen_values = [list(series.values) for series in gen_chart.series]
            ref_values = [list(series.values) for series in ref_chart.series]
            if all(value is None for values in gen_values for value in values):
                row['values_blank'] += 1
            elif gen_values == ref_values:
                row['values_same_as_reference'] += 1
            else:
                row['values_replaced'] += 1
        report.append(row)
    return report


def reference_data_match(loaded, template=TEMPLATE, tolerance=Decimal('0.5')):
    """Для каждой диаграммы шаблона со структурой План/Факт ищется строка Excel,
    период и филиал, воспроизводящие все точки. Возвращаются только номера."""
    index = {}
    for book in ('quantity', 'financial'):
        frame, provenance = loaded.frames[book], loaded.provenance[book]
        table = {}
        for position in frame.index:
            key = (frame.at[position, 'Наименование филиала'], frame.at[position, 'Период'])
            try:
                plan = Decimal(str(frame.at[position, 'Утверждено']))
                actual = Decimal(str(frame.at[position, 'Факт по 1С']))
            except Exception:
                continue
            table.setdefault(key, []).append((int(provenance.at[position, 'Строка Excel']), plan, actual))
        index[book] = table

    def close(left, right):
        return left is not None and right is not None and abs(Decimal(str(left)) - right) <= tolerance

    results = []
    deck = Presentation(template)
    for number, slide in enumerate(deck.slides, 1):
        if number < FIRST_TEMPLATE_SLIDE:
            continue
        for chart_index, shape in enumerate((s for s in slide.shapes if s.has_chart), 1):
            chart = shape.chart
            roles = series_roles(chart)
            categories = [category.label for category in chart.plots[0].categories]
            kind = classify_categories(categories)
            if roles is None or 'plan' not in roles or 'actual' not in roles or kind == 'unsupported':
                results.append({'slide': number, 'chart': chart_index, 'kind': kind, 'match': None, 'reason': 'ряды вне правил'})
                continue
            values = {role: list(series.values) for role, series in zip(roles, chart.series)}
            plan, actual = values['plan'], values['actual']
            best = None
            if kind == 'branch':
                candidates = {}
                for label, point_plan, point_actual in zip(categories, plan, actual):
                    branch = map_center(label, loaded.center_map)
                    if branch is None:
                        continue
                    for period in PERIOD_CANONICAL:
                        for book in index:
                            for row, excel_plan, excel_actual in index[book].get((branch, period), []):
                                if close(point_plan, excel_plan) and close(point_actual, excel_actual):
                                    key = (book, row, period)
                                    candidates[key] = candidates.get(key, 0) + 1
                for (book, row, period), hits in candidates.items():
                    if best is None or hits > best['hits']:
                        best = {'book': book, 'row': row, 'period': period, 'hits': hits}
            else:
                periods = [canonical_period(label) for label in categories]
                candidates = {}
                for period, point_plan, point_actual in zip(periods, plan, actual):
                    for book in index:
                        for branch in BRANCH_SHEETS:
                            for row, excel_plan, excel_actual in index[book].get((branch, period), []):
                                if close(point_plan, excel_plan) and close(point_actual, excel_actual):
                                    key = (book, row, branch)
                                    candidates[key] = candidates.get(key, 0) + 1
                for (book, row, branch), hits in candidates.items():
                    if best is None or hits > best['hits']:
                        best = {'book': book, 'row': row, 'branch': branch, 'hits': hits}
            mapped = sum(1 for label in categories if map_center(label, loaded.center_map)) if kind == 'branch' else len(categories)
            results.append({
                'slide': number, 'chart': chart_index, 'kind': kind, 'points': len(categories),
                'mapped': mapped, 'match': best,
                'reproduced': bool(best and best['hits'] == len(categories)),
            })
    return results


def run_variant_checks(loaded, keep_dir=OUTPUT_DIR):
    report = {'passed': [], 'diagnostics': []}
    before = source_hashes()
    unique_articles = [47, 54, 59]
    cases = [
        ('1x1x1', unique_articles[:1], ['9 месяцев'], BRANCH_SHEETS[:1]),
        ('3art_4per_6br', unique_articles, ['1 квартал', '2 квартал', '3 квартал', '9 месяцев'], BRANCH_SHEETS[:6]),
        ('3art_9m_18br', unique_articles, ['9 месяцев'], BRANCH_SHEETS[:18]),
        ('3art_9m_20br', unique_articles, ['9 месяцев'], BRANCH_SHEETS[:20]),
    ]
    with TemporaryDirectory(prefix='papagun-check-') as temporary:
        folder = Path(temporary)
        for name, articles, periods, branches in cases:
            path, diagnostics = generate_variant(
                loaded, articles, periods, branches,
                available_output(Path(keep_dir) / f'NL_variant_{name}.pptx'),
            )
            report['diagnostics'].extend(diagnostics)
            if path is None:
                raise AssertionError(f'Не собран проверочный файл {name}.')
            payloads = chart_payloads(path)
            if not payloads:
                raise AssertionError(f'В {name} нет диаграмм.')
            for article_row in articles:
                ok, reason = verify_fill_against_excel(
                    loaded, 'quantity', article_row, periods[0], branches[0],
                )
                if not ok:
                    raise AssertionError(reason or 'независимая сверка не прошла')
            report['passed'].append({
                'case': name,
                'slides': len(Presentation(path).slides),
                'charts': len(payloads),
                'series_points_checked': verify_series_rules(path),
                'articles': len(articles),
                'periods': len(periods),
                'branches': len(branches),
            })
        missing = generate_variant(
            loaded, [47], ['9 месяцев'], ['Нет такого филиала'], folder / 'missing.pptx',
        )[1]
        if not any(item.reason == 'филиал отсутствует в схеме загрузки' for item in missing):
            raise AssertionError('Отсутствующий филиал не диагностирован.')
        bad_period = generate_variant(loaded, [47], ['10 месяцев'], BRANCH_SHEETS[:1], folder / 'period.pptx')[1]
        if not any(item.reason == 'период отсутствует в схеме загрузки' for item in bad_period):
            raise AssertionError('Неизвестный период не диагностирован.')
        rows, reason = match_article('Количество процедур', loaded.articles['quantity'])
        if reason != 'статья неоднозначна' or not rows:
            raise AssertionError('Повтор названия статьи не распознан как неоднозначность.')
        for caption, expected in (('КБУ', 107), ('Утилизация м.о.', 128), ('Гепарин', 59), ('Транспортные расходы', 131)):
            rows, reason = match_article(caption, loaded.articles['quantity'])
            if reason or rows != (expected,):
                raise AssertionError(f'Подпись «{caption}» не привязана к ожидаемой строке.')
        if match_article('Шприцы', loaded.articles['quantity'])[1] != 'статья неоднозначна':
            raise AssertionError('«Шприцы» должны быть неоднозначны.')
        if match_article('РАСХОДЫ ЛС (количество)', loaded.articles['quantity'])[1] != 'подпись не является статьёй':
            raise AssertionError('Заголовок слайда принят за статью.')
        broken = LoadedBooks(
            frames={key: value.copy() for key, value in loaded.frames.items()},
            provenance=loaded.provenance,
            diagnostics=loaded.diagnostics,
            articles=loaded.articles,
            center_map=loaded.center_map,
            workbooks=loaded.workbooks,
        )
        target = broken.frames['quantity']
        selected = target.loc[
            target['Наименование филиала'].eq(BRANCH_SHEETS[0])
            & target['Период'].eq('9 месяцев')
            & loaded.provenance['quantity']['Строка Excel'].eq(47)
        ]
        if len(selected) != 1:
            raise AssertionError('Не найдена контрольная строка для нулевого плана.')
        target.at[selected.index[0], 'Утверждено'] = 0
        zero_path, zero_diag = generate_variant(
            broken, [47], ['9 месяцев'], BRANCH_SHEETS[:1], folder / 'zero.pptx',
        )
        if not any('нулевой план' in item.reason for item in zero_diag):
            raise AssertionError('Нулевой план не диагностирован.')
        if zero_path is None:
            raise AssertionError('Нулевой план не должен подменять значения и обязан оставлять диагностику.')
        zero_payloads = chart_payloads(zero_path)
        ratio_series = [values for _categories, series in zero_payloads for name, values in series if header_text(name) == 'столбец1']
        if not ratio_series or any(value is not None for values in ratio_series for value in values):
            raise AssertionError('Отношение при нулевом плане должно быть пустым, а не нулём.')
        target.at[selected.index[0], 'Утверждено'] = '#DIV/0!'
        error_diag = generate_variant(
            broken, [47], ['9 месяцев'], BRANCH_SHEETS[:1], folder / 'error.pptx',
        )[1]
        if not any('ошибка Excel' in item.reason for item in error_diag):
            raise AssertionError('Ошибка Excel не диагностирована.')
    assert_sources_unchanged(before)
    report['diagnostics'].extend(missing + bad_period + zero_diag + error_diag)
    report['passed'].extend([
        {'case': 'ambiguous_article'}, {'case': 'caption_matching'}, {'case': 'missing_branch'},
        {'case': 'unknown_period'}, {'case': 'zero_plan'}, {'case': 'excel_error'},
    ])
    return report


def main():
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    require_project_root()
    before = source_hashes()
    loaded = load_books()
    template_path, template_diag = generate_from_template(loaded)
    comparison = compare_with_reference(template_path)
    reference = reference_data_match(loaded)
    checks = run_variant_checks(loaded)
    assert_sources_unchanged(before)
    summary = {
        'quantity_rows': int(len(loaded.frames['quantity'])),
        'financial_rows': int(len(loaded.frames['financial'])),
        'template_file': template_path.name,
        'template_slides': len(Presentation(template_path).slides),
        'template_reasons': summarize_diagnostics(template_diag),
        'template_series_points_checked': verify_series_rules(template_path),
        'comparison': {
            'slides': len(comparison),
            'structure_equal_charts': sum(row['structure_equal'] for row in comparison),
            'charts_total': sum(row['charts'] for row in comparison),
            'values_replaced': sum(row['values_replaced'] for row in comparison),
            'values_blank': sum(row['values_blank'] for row in comparison),
            'values_same_as_reference': sum(row['values_same_as_reference'] for row in comparison),
            'texts_changed_slides': [row['slide'] for row in comparison if not row['texts_equal']],
        },
        'reference_reproducible_charts': [
            (item['slide'], item['chart'], item['match']['book'], item['match']['row'], item['match'].get('period') or item['match'].get('branch'))
            for item in reference if item.get('reproduced')
        ],
        'reference_charts_checked': sum(1 for item in reference if 'reproduced' in item),
        'reference_charts_reproduced': sum(1 for item in reference if item.get('reproduced')),
        'variant_cases': checks['passed'],
        'variant_reasons': summarize_diagnostics(checks['diagnostics']),
        'unfilled_template_slides': sorted(UNFILLED_TEMPLATE_SLIDES),
        'sources_sha256_unchanged': True,
    }
    for item in list(template_diag) + list(checks['diagnostics']):
        assert_no_values_in_text(item.reason)
    print('PASS: generate_presentation')
    for key, value in summary.items():
        print(f'{key}: {value}')
    return summary, template_diag, checks, comparison, reference


if __name__ == '__main__':
    main()
