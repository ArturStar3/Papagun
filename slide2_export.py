"""Local, editable analogue of source slide 2, populated from notebook DataFrames."""
from __future__ import annotations

from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
TEMPLATE = ROOT / 'Исходники/Презентация_NL_9 мес 2026.pptx'
CENTERS = ROOT / 'Исходники/Соответствие названий центров в презентации и эксель Nephrolife.xlsx'
MONTHS = ('январь', 'февраль', 'март', 'апрель', 'май', 'июнь',
          'июль', 'август', 'сентябрь', 'октябрь', 'ноябрь', 'декабрь')


def _center_key(text):
    key = re.sub(r'[^а-яёa-z0-9]', '', str(text).casefold().replace('\\n', ''))
    return key.replace('12мкрн', '12мкр')


def _number(value):
    if pd.isna(value) or isinstance(value, (str, bool)):
        raise ValueError('Показатель пациентов отсутствует или не является числом.')
    result = Decimal(str(value))
    if not result.is_finite() or result < 0 or result != result.to_integral_value():
        raise ValueError('Ожидается целое неотрицательное число пациентов.')
    return result


def _percent(value):
    return str((value * 100).quantize(Decimal('0.1'), rounding=ROUND_HALF_UP)).replace('.', ',') + '%'


def _integer(value):
    return f'{int(value):,}'.replace(',', ' ')


def patient_model(quantity_df, quantity_provenance, labels, *, month):
    """Experimental snapshot: source row 14, approved vs actual 1C; no month sums."""
    if month not in MONTHS:
        raise ValueError('Неизвестный месяц.')
    if not quantity_df.index.is_unique or not quantity_df.index.equals(quantity_provenance.index):
        raise ValueError('Индексы DataFrame и происхождения не совпадают или не уникальны.')
    if not quantity_provenance['Книга'].eq('quantity').all():
        raise ValueError('Ожидается происхождение количественного бюджета.')
    wb = load_workbook(CENTERS, read_only=True, data_only=True)
    try:
        mapping = {}
        for title, branch in list(wb.active.values)[1:]:
            if not title or not branch:
                continue
            # Explicit suffix difference between supplied directory and workbook.
            if branch == 'Усть-Каменогорск':
                branch = 'Усть-Каменогорск (NL UK)'
            key = _center_key(title)
            if key in mapping:
                raise ValueError('Повторяющееся название в справочнике центров.')
            mapping[key] = branch
    finally:
        wb.close()
    rows = []
    for label in labels:
        branch = mapping.get(_center_key(label))
        if branch is None:
            raise ValueError('Центр образца не найден в справочнике.')
        selected = quantity_df.loc[
            quantity_df['Наименование филиала'].eq(branch)
            & quantity_df['Месяц'].eq(month)
            & quantity_provenance['Строка Excel'].eq(14)
        ]
        if len(selected) != 1:
            raise ValueError('Для центра требуется ровно одна строка пациентов.')
        row = selected.iloc[0]
        origin = quantity_provenance.loc[selected.index[0]]
        if origin['Лист'] != branch or row['Наименование статей'].strip() != 'Количество пациентов':
            raise ValueError('Изменилась структура источника: строка 14 должна содержать пациентов.')
        plan, actual = _number(row['Утверждено']), _number(row['Факт по 1С'])
        if plan == 0:
            raise ValueError('Нулевой план: правило расчёта исполнения требует согласования.')
        rows.append(dict(label=label.strip(), branch=branch, plan=plan, actual=actual,
                         ratio=actual / plan, source_index=selected.index[0]))
    if len({r['branch'] for r in rows}) != len(rows):
        raise ValueError('Один филиал сопоставлен нескольким центрам.')
    # Original is sorted by execution; stable ordering resolves equal ratios.
    return sorted(rows, key=lambda r: r['ratio'], reverse=True)


def build_slide2(quantity_df, quantity_provenance, *, month='сентябрь',
                 output=ROOT / 'output/NL_пациенты_2026_слайд2.pptx'):
    """Use source styling; replace chart workbook, labels, title and all totals."""
    output = Path(output)
    if output.suffix.lower() != '.pptx':
        raise ValueError('Выходной файл должен иметь расширение .pptx.')
    if output.resolve().is_relative_to((ROOT / 'Исходники').resolve()):
        raise ValueError('Сохранение в каталог исходников запрещено.')
    if output.exists():
        raise FileExistsError('Результат уже существует; укажите новое имя файла.')
    deck = Presentation(TEMPLATE)
    slide = deck.slides[1]
    chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
    labels = [cat.label for cat in chart.plots[0].categories]
    model = patient_model(quantity_df, quantity_provenance, labels, month=month)
    # Drop unused slide relationships so unrelated slide data are not exported.
    for sld_id in list(deck.slides._sldIdLst):
        if deck.part.related_slide(sld_id.rId) is not slide:
            deck.part.drop_rel(sld_id.rId)
            deck.slides._sldIdLst.remove(sld_id)
    for rel in list(slide.part.rels.values()):
        if rel.reltype.endswith('/notesSlide'):
            slide.part.drop_rel(rel.rId)
    data = CategoryChartData(number_format='0')
    data.categories = [r['label'] for r in model]
    data.add_series('План', [int(r['plan']) for r in model])
    data.add_series('Факт', [int(r['actual']) for r in model])
    chart.replace_data(data)
    chart.plots[0].gap_width = 70
    chart.plots[0].overlap = -10
    # The source uses a third bar series for manually moved percentages. Replace
    # it with editable badges; remove per-point overrides tied to old values.
    for series in chart.series:
        for node in series._element.xpath('./c:dLbls/c:dLbl | ./c:dPt | ./c:extLst | ./c:spPr/a:effectLst'):
            node.getparent().remove(node)
    plan = sum((r['plan'] for r in model), Decimal(0))
    actual = sum((r['actual'] for r in model), Decimal(0))
    title, summary = [s for s in slide.shapes if s.has_text_frame]
    title.text_frame.paragraphs[0].runs[2].text = f'ПОКАЗАТЕЛИ: {month} 2026 г.'
    title.text_frame.paragraphs[0].runs[4].text = ' '
    pars = summary.text_frame.paragraphs
    pars[0].runs[0].text = f'{month.capitalize()} 2026: {len(model)} центров образца'
    pars[1].runs[1].text = _integer(plan) + ' пациентов'
    pars[2].runs[1].text = _integer(actual) + ' пациентов'
    pars[3].runs[1].text = _integer(actual - plan) + ' '
    pars[4].runs[1].text = _percent(actual / plan)
    color = RGBColor.from_string('FF0000' if actual < plan else '4472C4')
    pars[3].runs[1].font.color.rgb = color
    pars[4].runs[1].font.color.rgb = color
    # Plot boundaries follow the inner manual layout of the sample.
    plot_x, plot_w = 0.1943, 12.8881
    for i, row in enumerate(model):
        center = plot_x + plot_w * (i + 0.5) / len(model)
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                    Inches(center - .28), Inches(1.47),
                                    Inches(.56), Inches(.24))
        box.name = f'Исполнение {i+1}'
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(255, 255, 255)
        box.line.color.rgb = RGBColor(0, 0, 0)
        box.line.width = Pt(.8)
        frame = box.text_frame
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = _percent(row['ratio'])
        r.font.name, r.font.size, r.font.bold = 'Arial Narrow', Pt(12), True
        r.font.color.rgb = RGBColor.from_string('FF0000' if row['ratio'] < Decimal('.95') else '4472C4')
    # Remove unused sections/custom shows that may still refer to dropped slides.
    for node in deck.part._element.xpath('./p:custShowLst | ./p:extLst'):
        node.getparent().remove(node)
    output.parent.mkdir(parents=True, exist_ok=True)
    buffer = BytesIO()
    deck.save(buffer)
    with output.open('xb') as target:
        target.write(buffer.getvalue())
    return output
