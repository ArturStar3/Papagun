"""Local report worker. The calculation and slide functions come from excel_loading.ipynb."""

import json
import os
import sys
import tempfile
import re
from pathlib import Path


def asset(name):
    base = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent.parent))
    return base / name


NOTEBOOK = asset("excel_loading.ipynb")
CENTERS = asset("Соответствие названий центров в презентации и эксель Nephrolife.xlsx")
LOGO = asset("nephrolife_logo.jpg")
if not CENTERS.is_file():
    CENTERS = asset("Исходники/Соответствие названий центров в презентации и эксель Nephrolife.xlsx")
if not LOGO.is_file():
    LOGO = asset("resources/nephrolife_logo.jpg")
MONTHS = ("январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь")
BRANCHES = (
    "Актобе", "Актобе АМЦ", "Астана Биржан Сал", "Астана Сарыарка", "Атырау", "Атырау Нурсая",
    "Атырау Нурсая 2", "Караганда ЦАД", "Караганда Строителей", "Кентау", "Костанай", "Лисаковск",
    "Рудный", "Павлодар Дюсенова", "Павлодар Павлова", "Экибастуз", "Петропавловск",
    "Усть-Каменогорск (NL UK)", "Шымкент", "Шымкент Абая",
)


def checked_files(payload):
    quantity = Path(payload["quantity"])
    financial = Path(payload["financial"])
    if quantity.resolve() == financial.resolve() or any(p.suffix.lower() != ".xlsx" or not p.is_file() for p in (quantity, financial)):
        raise ValueError("Выберите две разные существующие книги .xlsx.")
    return {"quantity": quantity, "financial": financial}


def group_for(code):
    try:
        n = float(code)
    except (TypeError, ValueError):
        return "Прочие статьи"
    if n < 200:
        return "Статьи до 200"
    prefix = str(code).strip().split(".")[0]
    if prefix == "211": return "Лекарственные препараты"
    if prefix == "212": return "Изделия медицинского назначения"
    if prefix == "213": return "Расходные медицинские материалы"
    if int(prefix) >= 214: return "Статьи от 214 и выше"
    return "Прочие статьи"


def inspect(payload):
    from openpyxl import load_workbook

    files = checked_files(payload)
    items = {}
    available = None
    month_sets = []
    for book, path in files.items():
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = [name for name in BRANCHES if name in wb.sheetnames]
            if not sheets:
                raise ValueError(f"В книге {book} не найдены листы филиалов NL.")
            available = set(sheets) if available is None else available & set(sheets)
            # A month is selectable only when the source has actual values for it.
            month_columns = (11, 15, 19, 23, 27, 31, 35, 39, 43, 47, 51, 55)
            valid_months = set()
            for branch in sheets:
                ws = wb[branch]
                for row in ws.iter_rows(min_row=9, max_row=141 if book == "quantity" else 193,
                                        min_col=1, max_col=55, values_only=True):
                    for month, column in zip(MONTHS, month_columns):
                        value = row[column - 1]
                        if isinstance(value, (int, float)) and not isinstance(value, bool):
                            valid_months.add(month)
                    if len(valid_months) == len(MONTHS):
                        break
            month_sets.append(valid_months)
            if book != "quantity":
                continue
            limit = 141
            for branch in sheets:
                ws = wb[branch]
                if str(ws["B6"].value).strip().casefold() != "код статей":
                    raise ValueError(f"Не найден заголовок кодов: {book}, {branch}, B6.")
                for row in ws.iter_rows(min_row=9, max_row=limit, min_col=2, max_col=3, values_only=True):
                    code, name = row
                    if code is None or name is None or not str(code).strip() or not str(name).strip():
                        continue
                    code, name = str(code).strip(), str(name).strip()
                    if not re.fullmatch(r"\d+(?:\.\d+)*\.?", code):
                        continue
                    key = (code, name)
                    items[key] = {"id": "\u001f".join(key), "code": code, "name": name, "group": group_for(code)}
        finally:
            wb.close()
    ordered = sorted(items.values(), key=lambda x: (x["group"], x["code"], x["name"]))
    branches = [b for b in BRANCHES if b in available]
    # The reference report contains 19 centers; the twentieth source sheet lacks
    # required September financial facts and has no approved missing-value rule.
    report_branches = [b for b in branches if b != "Атырау Нурсая 2"]
    return {"months": MONTHS, "availableMonths": [m for m in MONTHS[:9] if all(m in s for s in month_sets)],
            "branches": branches, "availableBranches": report_branches, "articles": ordered}


def notebook_namespace(files, branches, months):
    import pandas as pd

    cells = json.loads(NOTEBOOK.read_text(encoding="utf-8"))["cells"]
    ns = {"__name__": "papagun_notebook"}
    for index in (3, 5, 12):
        source = "".join(cells[index]["source"])
        if index == 5:
            source = source.replace("ROOT = Path.cwd()", "ROOT = Path.cwd()")
            source = source.replace("if not (ROOT / 'excel_loading.ipynb').is_file():", "if False:")
        exec(compile(source, str(NOTEBOOK), "exec"), ns)
    ns["CENTERS"] = CENTERS
    ns["LOGO"] = LOGO
    ns["BRANCH_SHEETS"] = list(branches)
    ns["excel_files"] = files
    ns["workbooks"] = {
        key: pd.read_excel(path, sheet_name=list(branches), header=None, dtype=object,
                           keep_default_na=False, engine="openpyxl")
        for key, path in files.items()
    }
    ns["REPORT_PERIODS"] = tuple(months)
    source = "".join(cells[13]["source"])
    source = source[:source.index("quantity_df, quantity_provenance")]
    exec(compile(source, str(NOTEBOOK), "exec"), ns)
    for index in (17, 19, 21):
        exec(compile("".join(cells[index]["source"]), str(NOTEBOOK), "exec"), ns)
    original_center_map = ns["slide5_center_map"]
    def current_center_map():
        mapping = original_center_map()
        mapping.update({ns["_center_key"](branch): branch for branch in branches})
        return mapping
    ns["slide5_center_map"] = current_center_map
    return ns


def financial_selector(frame, code, quantity_name, branches, months):
    """Use the quantity key; accept a different finance label only for an unambiguous code."""
    codes = frame["Код статей"].map(lambda value: value.strip() if isinstance(value, str) else value)
    matches = frame.loc[codes.eq(code)
                        & frame["Наименование филиала"].isin(branches)
                        & frame["Месяц"].isin(months)]
    names = set(matches["Наименование статей"].dropna())
    if not names:
        raise ValueError(f"Код {code} отсутствует в финансовой книге.")
    if quantity_name in names:
        return (code, quantity_name)
    if len(names) == 1:
        return code
    raise ValueError(f"Код {code} имеет несколько финансовых наименований; сопоставление с «{quantity_name}» неоднозначно.")


def article_groups(items):
    """Keep categories independent; use a three-item last page for odd counts."""
    count = len(items)
    blocks = []
    index = 0
    while count - index > 3:
        blocks.append(items[index:index + 2])
        index += 2
    if index < count:
        blocks.append(items[index:])
    return blocks


def period_caption(months):
    indexes = [MONTHS.index(month) for month in months]
    if indexes == list(range(indexes[0], indexes[-1] + 1)):
        return months[0] if len(months) == 1 else f"{months[0]}–{months[-1]}"
    return f"{len(months)} выбранных месяцев"


def presentation_labels(branches):
    from openpyxl import load_workbook

    workbook = load_workbook(CENTERS, read_only=True, data_only=True)
    try:
        labels = {"Усть-Каменогорск (NL UK)": "Усть-Каменогорск"}
        labels.update({str(branch).strip(): str(label).strip() for label, branch in list(workbook.active.values)[1:]
                       if label and branch and str(branch).strip() != "Усть-Каменогорск"})
    finally:
        workbook.close()
    return tuple(labels.get(branch, branch) for branch in branches)


def combined_model(ns, frame, provenance, branches, months, book, item):
    from decimal import Decimal

    components = item["components"]
    selectors = item["selectors"][book]
    raw = ns["_drug_models_by_key"](
        frame, provenance, list(branches), months=months, book=book,
        first_metric="Утверждено" if book == "quantity" else "Факт по норме",
        divisor=Decimal(1 if book == "quantity" else 1000), articles=selectors,
    )
    if len(components) == 1:
        return raw[0]
    result = []
    for members in zip(*raw):
        if len({(row["label"], row["branch"]) for row in members}) != 1:
            raise ValueError("Компоненты группы относятся к разным филиалам.")
        row = dict(label=members[0]["label"], branch=members[0]["branch"],
                   article=item["title"], article_code=item["key"],
                   component_keys=tuple((r["article_code"], r["article"]) for r in members),
                   source_indices=tuple(i for r in members for i in r["source_indices"]))
        for metric in ("plan", "actual", "norm"):
            row[metric] = sum((r[metric] for r in members), Decimal(0))
        row["difference"] = row["norm"] - row["actual"]
        result.append(row)
    return result


def quarter_models(ns, frame, provenance, branches, months, book, items):
    from decimal import Decimal

    # Group the chosen months by calendar quarter; never create a zero for an unchosen quarter.
    selected_months = set(months)
    periods = [(f"{q + 1} квартал", tuple(m for m in MONTHS[q * 3:q * 3 + 3] if m in selected_months))
               for q in range(4)]
    periods = [(label, part) for label, part in periods if part]
    result = []
    for item in items:
        rows = []
        for label, period_months in periods:
            centers = combined_model(ns, frame, provenance, branches, period_months, book, item)
            norm = sum((r["norm"] for r in centers), Decimal(0))
            actual = sum((r["actual"] for r in centers), Decimal(0))
            rows.append(dict(label=label, article=item["title"], article_code=item["key"],
                             plan=norm, norm=norm, actual=actual, difference=norm - actual))
        norm = sum((r["norm"] for r in rows), Decimal(0))
        actual = sum((r["actual"] for r in rows), Decimal(0))
        rows.append(dict(label="9 месяцев" if tuple(months) == MONTHS[:9] else "Выбранный период", article=item["title"], article_code=item["key"],
                         plan=norm, norm=norm, actual=actual, difference=norm - actual))
        result.append(rows)
    return result


def render_quarters(ns, models, title):
    deck = ns["new_presentation"]()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    ns["add_header"](slide, title, width=11.8)
    count = len(models)
    for index, rows in enumerate(models):
        if count == 4:
            top, height = .65 + index * 1.67, 1.63
        elif count == 3:
            top, height = .65 + index * 2.16, 2.04
        else:
            top, height = ((.65, 3.10) if index == 0 else (3.88, 3.12))
        series = [(name, [r[key] for r in rows], color) for name, key, color in
                  zip(("Факт по норме", "Факт по 1С", "отк:"), ("norm", "actual", "difference"),
                      (ns["STYLE"]["teal"], ns["STYLE"]["gray"], ns["STYLE"]["gray"]))]
        chart = ns["numeric_chart"](slide, [r["label"] for r in rows], series,
                                     .15, top, 12.99, height, fmt="#,##0", compact_base=len(rows),
                                     drug=True, total=True)
        if count == 4:
            ns["set_plot_layout"](chart.chart, x=.01, y=.30, w=.98, h=.55)
        ns["text_box"](slide, rows[0]["article"], .25, top + .025, 12.79, .22,
                       size=9, bold=True, color="0070C0", name="quarter_article")
        for j, row in enumerate(rows):
            center = .15 + 12.99 * (.01 + (j + .5) * .98 / len(rows))
            percent, color = ns["quarter_execution_info"](row["norm"], row["actual"])
            panel = ns["text_box"](slide, percent, center - .325, top + (.22 if count == 4 else .12),
                                   .65, .18, size=9, bold=True, color=color, fill=ns["STYLE"]["gray"])
            panel.line.color.rgb = ns["RGBColor"](0, 0, 0)
    return ns["presentation_bytes"](deck)


def statistic_slide(ns, frame, provenance, branches, labels, months, row_number, code, title, forms):
    from decimal import Decimal

    period = (months[-1],) if row_number != 15 else tuple(months)
    model = []
    for branch, label in zip(branches, labels):
        selected = frame.loc[frame["Наименование филиала"].eq(branch)
                             & frame["Месяц"].isin(period)
                             & provenance["Строка Excel"].eq(row_number)]
        if (len(selected) != len(period) or set(selected["Месяц"]) != set(period)
                or selected["Месяц"].duplicated().any()
                or not selected["Код статей"].map(lambda value: str(value).strip()).eq(code).all()
                or not provenance.loc[selected.index, "Лист"].eq(branch).all()):
            raise ValueError(f"Нет однозначных данных показателя «{title}» для филиала «{branch}».")
        number = ns["_procedure_number"] if row_number == 15 else ns["_number"]
        plan = sum((number(value) for value in selected["Утверждено"]), Decimal(0))
        actual = sum((number(value) for value in selected["Факт по 1С"]), Decimal(0))
        model.append(dict(label=label, plan=plan, actual=actual))
    deck = ns["new_presentation"]()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    ns["add_header"](slide, f"{title}: {period_caption(period)} 2026 г.", width=11.8)
    plan = sum((r["plan"] for r in model), Decimal(0))
    actual = sum((r["actual"] for r in model), Decimal(0))
    ns["add_summary"](slide, [(f"{len(model)} центров", "000000"),
                               ("План: " + ns["_counted"](plan, forms), "000000"),
                               ("Факт: " + ns["_counted"](actual, forms), "000000"),
                               ("Отклонение: " + ns["_counted"](actual - plan, forms), "000000")])
    ns["numeric_chart"](slide, [r["label"] for r in model],
                         [("План", [r["plan"] for r in model], ns["STYLE"]["teal"]),
                          ("Факт", [r["actual"] for r in model], ns["STYLE"]["gray"])],
                         0, 1.85, 13 + 1 / 3, 5.5, fmt="General", compact_base=len(model))
    return ns["presentation_bytes"](deck)


def generate(payload):
    import hashlib
    from io import BytesIO
    from zipfile import ZipFile
    from pptx import Presentation

    files = checked_files(payload)
    months = tuple(m for m in MONTHS if m in payload["months"])
    branches = tuple(b for b in BRANCHES if b in payload["branches"])
    articles = payload.get("articles", [])
    custom_groups = payload.get("groups", [])
    if not months or not branches:
        raise ValueError("Выберите хотя бы один месяц и филиал.")
    if len(months) != len(set(payload["months"])) or len(branches) != len(set(payload["branches"])):
        raise ValueError("Список месяцев или филиалов содержит неизвестные или повторяющиеся значения.")
    available_catalog = inspect(payload)
    if any(m not in available_catalog["availableMonths"] for m in months) or any(b not in available_catalog["availableBranches"] for b in branches):
        raise ValueError("Выбран месяц или филиал без полного набора данных для этого отчёта.")
    output = Path(payload["output"])
    if output.suffix.lower() != ".pptx" or output.resolve() in (p.resolve() for p in files.values()):
        raise ValueError("Укажите выходной файл .pptx, отличный от исходных книг.")
    output.parent.mkdir(parents=True, exist_ok=True)
    before = {key: hashlib.sha256(path.read_bytes()).digest() for key, path in files.items()}
    ns = notebook_namespace(files, branches, months)
    frames = {}
    for book in files:
        frames[book] = ns["build_branches_dataframe"](book, periods=months)[:2]
    catalog = {item["id"]: item for item in available_catalog["articles"]}
    used = set()

    def parse(raw):
        if raw not in catalog:
            raise ValueError("Выбрана статья, которой нет в количественной книге.")
        parts = raw.split("\u001f")
        if len(parts) != 2 or not all(parts):
            raise ValueError("Некорректная статья в выборе.")
        if raw in used:
            raise ValueError("Статья используется повторно в выборе или группе.")
        used.add(raw)
        if catalog[raw]["group"] not in ("Лекарственные препараты", "Изделия медицинского назначения", "Расходные медицинские материалы"):
            raise ValueError("Отдельные слайды строятся только для лекарств, медицинских изделий и расходных материалов.")
        if catalog[raw]["code"] in ("211", "212", "213"):
            raise ValueError("Выберите исходные статьи, а не итоговую строку раздела.")
        return tuple(parts), catalog[raw]["group"]

    items = []
    for raw in articles:
        pair, category = parse(raw)
        items.append(dict(title=pair[1], key=pair[0], category=category, components=[pair]))
    for entry in custom_groups:
        title = str(entry.get("title", "")).strip()
        members = entry.get("articles", [])
        if not title or len(members) < 2 or len(members) != len(set(members)):
            raise ValueError("Группе нужны заголовок и хотя бы две разные статьи.")
        parsed = [parse(raw) for raw in members]
        categories = {category for _, category in parsed}
        if len(categories) != 1:
            raise ValueError("В одной группе можно объединять статьи только из одного раздела.")
        components = [pair for pair, _ in parsed]
        items.append(dict(title=title, key=" + ".join(code for code, _ in components),
                          category=categories.pop(), components=components))
    financial_used = []
    for item in items:
        item["selectors"] = {"quantity": item["components"],
                             "financial": [financial_selector(frames["financial"][0], code, name, branches, months)
                                           for code, name in item["components"]]}
        financial_used.extend(item["selectors"]["financial"])
    if len(set(financial_used)) != len(financial_used):
        raise ValueError("Несколько выбранных количественных позиций соответствуют одной финансовой статье. Сопоставление неоднозначно.")
    payloads = []
    quantity, quantity_provenance = frames["quantity"]
    labels = presentation_labels(branches)
    for row, code, title, forms in ((13, "2", "АППАРАТЫ", ("аппарат", "аппарата", "аппаратов")),
                                    (14, "3", "ПАЦИЕНТЫ", ("пациент", "пациента", "пациентов")),
                                    (15, "4", "ПРОЦЕДУРЫ", ("процедура", "процедуры", "процедур"))):
        payloads.append(statistic_slide(ns, quantity, quantity_provenance, branches, labels, months, row, code, title, forms))
    category_titles = {"Лекарственные препараты": "ЛС", "Изделия медицинского назначения": "ИМН",
                       "Расходные медицинские материалы": "РАСХОДНЫЕ МАТЕРИАЛЫ"}
    for category in category_titles:
        for block in article_groups([item for item in items if item["category"] == category]):
            for book in ("quantity", "financial"):
                frame, provenance = frames[book]
                models = [combined_model(ns, frame, provenance, labels, months, book, item) for item in block]
                selectors = [(item["key"], rows[0]["article"]) for item, rows in zip(block, models)]
                heading = ("РАСХОДЫ " + category_titles[category] + ": " + period_caption(months) + " 2026 г."
                           + (" (количество)" if book == "quantity" else " (тыс. тенге)"))
                payloads.append(ns["render_drug_models"](
                    models, heading=heading, series_names=("План", "Факт", "отк:") if book == "quantity"
                    else ("Факт по норме", "Факт по 1С", "отк:"), financial=book == "financial",
                    articles=selectors))
            models = quarter_models(ns, quantity, quantity_provenance, labels, months, "quantity", block)
            payloads.append(render_quarters(ns, models, "РАСХОДЫ " + category_titles[category] + ": по кварталам 2026 г. (количество)"))

    financial, financial_provenance = frames["financial"]
    expense_items = [dict(title=label, key=" + ".join(codes), components=[(code, "") for code in codes],
                          selectors={"financial": list(codes)}) for label, codes in (
                              ("Транспортные расходы", ("223.01", "223.02", "223.03", "223.04")),
                              ("Расходы на питание", ("222.00", "222.01")))]
    expense_models = [combined_model(ns, financial, financial_provenance, labels, months, "financial", item)
                      for item in expense_items]
    payloads.append(ns["render_drug_models"](
        expense_models, heading="РАСХОДЫ: " + period_caption(months) + " 2026 г. (тыс. тенге)",
        series_names=("Факт по норме", "Факт по 1С", "отк:"), financial=True,
        articles=[(item["key"], item["title"]) for item in expense_items]))
    quarter_items = [dict(title=label, key=" + ".join(codes), components=[(code, "") for code in codes],
                          selectors={"financial": list(codes)}) for label, codes in (
                              ("Расходы на транспорт", ("223.01", "223.02")),
                              ("Возмещение транспортных расходов пациентам", ("223.04",)),
                              ("Расходы на питание", ("222.00",)),
                              ("Возмещение расходов на питание", ("222.01",)))]
    payloads.append(render_quarters(ns, quarter_models(ns, financial, financial_provenance, labels, months,
                                                       "financial", quarter_items),
                                    "ФИНАНСОВЫЕ ПОКАЗАТЕЛИ: по кварталам 2026 г. (тыс. тенге)"))
    result = ns["combine_created_slides"](payloads)
    with ZipFile(BytesIO(result)) as archive:
        if archive.testzip() is not None:
            raise ValueError("Созданный PPTX повреждён.")
    if len(Presentation(BytesIO(result)).slides) != len(payloads):
        raise ValueError("Количество слайдов не совпало с выбранными статьями.")
    for key, path in files.items():
        if hashlib.sha256(path.read_bytes()).digest() != before[key]:
            raise ValueError("Исходная книга изменилась во время обработки.")
    fd, temp_path = tempfile.mkstemp(suffix=".pptx", dir=output.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(result)
        os.replace(temp_path, output)
    finally:
        if os.path.exists(temp_path):
            os.unlink(temp_path)
    return {"output": str(output), "slides": len(payloads), "articles": len(used)}


if __name__ == "__main__":
    sys.stdin.reconfigure(encoding="utf-8")
    sys.stdout.reconfigure(encoding="utf-8")
    try:
        request = json.load(sys.stdin)
        response = inspect(request) if sys.argv[1] == "inspect" else generate(request)
        print(json.dumps(response, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False))
        sys.exit(1)
